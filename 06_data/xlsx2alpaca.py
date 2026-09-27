# -*- coding: utf-8 -*-
"""专家经验话术汇总.xlsx  →  LLaMA-Factory Alpaca 训练集

用法：
    pip install openpyxl
    python 06_data/xlsx2alpaca.py 专家经验话术汇总-板带.xlsx

产出：
    06_data/huashu.jsonl       训练集（Alpaca 格式）
    05_kb/facts_seed.tsv       事实库种子（从话术里抠出的数字/型号/认证）
    06_data/convert_report.txt 转换报告 + 数据健康度

── 两个关键设计，别随手改掉 ──

1) 【可用事实】槽位（oracle RAG）
   把 output 里出现的数字、型号、认证名抠出来，塞回 input 的【可用事实】里。
   这样模型学到的是「用给你的事实」，而不是「在这个位置编一个数字」。
   等阶段 1 的知识库建好，把 oracle 事实换成检索结果，数据格式不用改。
   → 关了这个开关（--no-oracle），训出来的模型会编造数字。见 方案v2 第三节。

2) 问题类型归一
   原表里有 5 处不一致（价格转价值→价格、客诉转提升品牌/客诉转品牌价值→客诉）。
   这里统一修掉。不修的话模型会把同一概念学成多个标签。
"""
import argparse
import json
import os
import re
import sys
from collections import Counter

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit('❌ 缺依赖，先跑：pip install openpyxl')

# ───────────────────────── 配置 ─────────────────────────

# sheet 名 → 目录。汇总表不含正文，跳过。
SHEET2CATALOG = {
    '专家经验回复话术':     '解决方案话术',
    '如何挖掘客户痛点话术': '如何挖掘客户痛点',
    '价格转价值话术':       '价格转价值话术',
    '客诉转品牌提升话术':   '客诉转品牌提升话术',
}
SKIP_SHEETS = {'专家经验话术汇总'}

# 列名 → 内部字段。用「包含」匹配，容忍表头里的括号和说明文字。
COL_ALIASES = [
    ('事业部',      ['事业部']),
    ('一级行业',    ['一级行业']),
    ('二级行业',    ['二级行业']),
    ('三级行业',    ['三级行业', '三级行']),
    ('产品大类',    ['EDM产品大类', '产品大类']),
    ('问题类型',    ['问题类型']),
    ('适用背景',    ['适用背景']),
    ('问题',        ['问题']),          # 放在「问题类型」之后，避免抢匹配
    ('思路措施',    ['思路及措施', '思路和措施']),
    ('标准话术',    ['标准沟通话术', '沟通话术']),
    ('备注',        ['备注']),
]

# 问题类型归一：左边是表里出现过的写法，右边是冻结后的标签
NORMALIZE = {
    '价格转价值':     '价格',
    '客诉转提升品牌': '客诉',
    '客诉转品牌价值': '客诉',
    '客诉转品牌提升': '客诉',
    '服务/物流':      '物流/服务',
    '服务、物流':     '物流/服务',
    '物流/服务':      '物流/服务',
}

# 目录 → 合法的问题类型（方案v2 第二节的冻结标签集）
LEGAL = {
    '解决方案话术':       ['价格', '交期', '客诉', '物流/服务', '新产品增值'],
    '如何挖掘客户痛点':   ['商务痛点', '生产痛点', '研发/技术痛点', '个人需求痛点'],
    '价格转价值话术':     ['价格'],
    '客诉转品牌提升话术': ['客诉'],
}

# 事实项抽取规则。宁可多抓（人工删）也不要漏抓（漏了模型就会编）。
# ⚠️ 每组内部**长的写在前面**——正则的 | 是先到先得，
#    把 K 写在 KPCS 前面，「500Kpcs」会被截成「500K」。
FACT_PATTERNS = [
    (r'\d+(?:\.\d+)?\s*(?:元/吨|元/㎡|万元|亿|万|千|元|吨|台|道|个|家|年|月|天|小时|分钟|次)', '数量'),
    (r'\d+(?:\.\d+)?\s*(?:KPCS|Kpcs|kpcs|kPCS|PPM|ppm|pcs|％|%|℃|度|K|k)',                   '指标'),
    (r'\bC\d{3,4}[A-Z]*\b',                                                                  '牌号'),
    (r'\b(?:IATF|ISO|TS)\s?\d{4,5}\b',                                                       '认证'),
    (r'\b(?:MES|SCADA|VAVE|CPO|EDM|Ra)\b',                                                   '术语'),
]

INSTRUCTION = """你是博威合金的资深销售专家。根据给定的业务场景和可用事实，输出两部分内容：

【解决问题的思路及措施】
按「思路N：…」+「措施N：…」的结构展开，每条思路配一条对应措施。

【标准沟通话术】
定性定量——定量必须用具体数据和案例说话。按场景或按「一、二、三」分段组织。

硬性要求：所有数字、产品牌号、认证名称、客户与竞品名，只能取自【可用事实】。
【可用事实】里没有的，写「（待补数据）」占位，不得自行编造。"""


# ───────────────────────── 工具 ─────────────────────────

def norm(s):
    return re.sub(r'\s+', '', str(s or ''))


def map_header(row):
    """表头行 → {内部字段: 列下标}。按 COL_ALIASES 顺序匹配，先到先得。"""
    used, out = set(), {}
    for field, pats in COL_ALIASES:
        for i, cell in enumerate(row):
            if i in used:
                continue
            h = norm(cell)
            if h and any(p in h for p in pats):
                out[field] = i
                used.add(i)
                break
    return out


def find_header(ws, scan=6):
    """表头不一定在第 1 行（有的 sheet 有标题行），前几行里找命中最多的那行。"""
    best, best_n, best_i = {}, 0, 1
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=scan, values_only=True), 1):
        m = map_header(row)
        if len(m) > best_n:
            best, best_n, best_i = m, len(m), i
    return best, best_i


def extract_facts(text):
    """抠出事实项，去重保序。"""
    seen, out = set(), []
    for pat, kind in FACT_PATTERNS:
        for m in re.finditer(pat, text):
            v = m.group().strip()
            if v not in seen:
                seen.add(v)
                out.append((v, kind))
    return out


# ───────────────────────── 主流程 ─────────────────────────

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('--out',   default='06_data/huashu.jsonl')
    ap.add_argument('--kb',    default='05_kb/facts_seed.tsv')
    ap.add_argument('--report', default='06_data/convert_report.txt')
    ap.add_argument('--no-oracle', action='store_true',
                    help='不往 input 里塞【可用事实】——⚠️ 会训出会编数字的模型')
    a = ap.parse_args()

    if not os.path.exists(a.xlsx):
        sys.exit(f'❌ 找不到 {a.xlsx}')

    wb = load_workbook(a.xlsx, data_only=True)   # data_only：取公式的计算结果
    rows, log, kb = [], [], []
    stat = Counter()
    fixed, illegal = [], []

    for sname in wb.sheetnames:
        if norm(sname) in {norm(s) for s in SKIP_SHEETS}:
            log.append(f'[跳过] {sname}（汇总表，无正文）')
            continue
        catalog = next((v for k, v in SHEET2CATALOG.items() if norm(k) == norm(sname)), None)
        if catalog is None:
            log.append(f'[跳过] {sname}（未知 sheet，如需纳入请加到 SHEET2CATALOG）')
            continue

        ws = wb[sname]
        cols, hrow = find_header(ws)
        log.append(f'\n[处理] {sname} → 目录「{catalog}」，表头在第 {hrow} 行')
        log.append(f'        识别到的列：{", ".join(cols)}')
        for need in ('思路措施', '标准话术'):
            if need not in cols:
                log.append(f'        ⚠️  没找到「{need}」列，这个 sheet 出不了训练样本')

        carry = {}   # 合并单元格造成的空值，用上一行的值补（事业部/行业/目录常被合并）
        for ri, row in enumerate(ws.iter_rows(min_row=hrow + 1, values_only=True), hrow + 1):
            get = lambda f: str(row[cols[f]] or '').strip() if f in cols and cols[f] < len(row) else ''

            for f in ('事业部', '一级行业', '二级行业', '三级行业', '产品大类', '问题类型'):
                v = get(f)
                if v:
                    carry[f] = v
            rec = {f: carry.get(f, '') for f in
                   ('事业部', '一级行业', '二级行业', '三级行业', '产品大类', '问题类型')}
            rec['适用背景'] = get('适用背景')
            rec['问题']     = get('问题')
            sol, talk       = get('思路措施'), get('标准话术')

            if not (sol or talk):
                stat['空行'] += 1
                continue

            # 问题类型归一
            raw = norm(rec['问题类型'])
            if raw in NORMALIZE and NORMALIZE[raw] != rec['问题类型']:
                fixed.append(f'{sname} 第{ri}行: 「{rec["问题类型"]}」→「{NORMALIZE[raw]}」')
                rec['问题类型'] = NORMALIZE[raw]
            if rec['问题类型'] and rec['问题类型'] not in LEGAL[catalog]:
                illegal.append(f'{sname} 第{ri}行: 「{rec["问题类型"]}」不在目录「{catalog}」的合法集内')

            # 组装 output
            parts = []
            if sol:
                parts.append('【解决问题的思路及措施】\n' + sol)
            if talk:
                parts.append('【标准沟通话术】\n' + talk)
            output = '\n\n'.join(parts)

            # 组装 input
            lines = [f'目录：{catalog}']
            for f in ('事业部', '一级行业', '二级行业', '三级行业', '产品大类',
                      '问题类型', '适用背景', '问题'):
                if rec.get(f):
                    lines.append(f'{f}：{rec[f]}')

            facts = extract_facts(output)
            for v, kind in facts:
                kb.append((v, kind, f'{sname}#{ri}'))
            if not a.no_oracle:
                lines.append('\n【可用事实】')
                lines.extend(f'- [{k}] {v}' for v, k in facts) if facts \
                    else lines.append('- （无）')

            rows.append({'instruction': INSTRUCTION,
                         'input': '\n'.join(lines),
                         'output': output})
            stat['样本'] += 1
            stat[f'类型:{rec["问题类型"] or "（空）"}'] += 1
            stat['事实项'] += len(facts)
            if not facts:
                stat['零事实样本'] += 1

    if not rows:
        sys.exit('❌ 一条样本都没出来。检查 sheet 名和列名是否和 SHEET2CATALOG / COL_ALIASES 对得上。')

    # ── 写文件 ──
    for p in (a.out, a.kb, a.report):
        os.makedirs(os.path.dirname(p) or '.', exist_ok=True)

    with open(a.out, 'w', encoding='utf-8') as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + '\n')

    with open(a.kb, 'w', encoding='utf-8') as f:
        f.write('事实项\t类别\t出处\t是否核实\n')
        for v, kind, src in sorted(set(kb)):
            f.write(f'{v}\t{kind}\t{src}\t待核实\n')

    # ── 健康度 ──
    lens = sorted(len(r['input']) + len(r['output']) for r in rows)
    n = len(lens)
    health = [
        '',
        '═══════════ 数据健康度 ═══════════',
        f'样本数              {stat["样本"]}',
        f'跳过空行            {stat["空行"]}',
        f'事实项总数          {stat["事实项"]}（去重后 {len(set(kb))}）',
        f'零事实样本          {stat["零事实样本"]}  ← 这些样本的话术里没有任何具体数字',
        '',
        f'字符长度  最短 {lens[0]}  中位 {lens[n//2]}  最长 {lens[-1]}',
        f'          → cutoff_len 至少要 {int(lens[-1] / 1.4 * 1.3 // 256 + 1) * 256} '
        f'（按中文 1.4 字/token、留 30% 余量估）',
        '',
        '问题类型分布：',
    ]
    for k, v in sorted(stat.items()):
        if k.startswith('类型:'):
            health.append(f'  {k[3:]:<16} {v}')

    warn = []
    if stat['零事实样本'] > stat['样本'] * 0.3:
        warn.append('⚠️  超过 30% 的样本没有具体数字——要么抽取规则漏了，'
                    '要么这些话术本身就不「定性定量」')
    if lens[-1] / 1.4 > 3500:
        warn.append(f'⚠️  最长样本约 {int(lens[-1]/1.4)} token，'
                    'cutoff_len 设小了会被静默截断，模型会学成「说一半就停」')
    if illegal:
        warn.append(f'⚠️  {len(illegal)} 处问题类型不在合法集内，见下方明细')

    report = '\n'.join(log + health + [''] +
                       (['═══════════ 已自动修正的标签 ═══════════'] + fixed if fixed else []) +
                       ([''] + ['═══════════ 非法标签（需人工处理）═══════════'] + illegal if illegal else []) +
                       [''] + (['═══════════ 警告 ═══════════'] + warn if warn else ['✅ 无警告']))
    open(a.report, 'w', encoding='utf-8').write(report + '\n')

    print(report)
    print(f'\n→ 训练集    {a.out}       （{stat["样本"]} 条）')
    print(f'→ 事实种子  {a.kb}   （{len(set(kb))} 项，**需人工核实出处**）')
    print(f'→ 报告      {a.report}')
    if a.no_oracle:
        print('\n⚠️  你用了 --no-oracle：input 里没有【可用事实】，'
              '训出来的模型会编造数字。除非你明确想对比，否则别这么用。')


if __name__ == '__main__':
    main()
